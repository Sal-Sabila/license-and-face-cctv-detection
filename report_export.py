# ============================================================
# report_export.py
# Export Excel & PDF untuk Sistem Monitoring CCTV
# ============================================================

import io
import os
from datetime import datetime
from urllib.parse import unquote
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.styles import (
    Font,
    PatternFill,
    Border,
    Side,
    Alignment,
)
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
    Image,
    PageBreak,
)
from reportlab.lib.utils import ImageReader


# ============================================================
# KONFIGURASI
# ============================================================

BASE_DIR = os.path.dirname(os.path.abspath(__file__))

EXCEL_MIMETYPE = (
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
)

PDF_MIMETYPE = "application/pdf"


# ============================================================
# WARNA EXCEL
# ============================================================

DARK = "1F2937"
HEADER = "2563EB"
LIGHT_BLUE = "DBEAFE"
LIGHT_GREEN = "DCFCE7"
LIGHT_YELLOW = "FEF3C7"
LIGHT_RED = "FEE2E2"
LIGHT_GRAY = "F3F4F6"
WHITE = "FFFFFF"
TEXT = "111827"
BORDER_COLOR = "D1D5DB"


# ============================================================
# HELPER UMUM
# ============================================================

def _value(data, *keys, default=""):
    """
    Mengambil nilai dari dictionary berdasarkan beberapa kemungkinan key.
    Membantu agar export tetap aman jika nama field sedikit berbeda.
    """
    if not isinstance(data, dict):
        return default

    for key in keys:
        if key in data and data[key] is not None:
            return data[key]

    return default


def _safe_text(value):
    if value is None:
        return "-"
    return str(value)


def _safe_number(value, default=0):
    try:
        if value is None or value == "":
            return default
        return float(value)
    except (TypeError, ValueError):
        return default


def _confidence(value):
    """
    Normalisasi confidence.

    Database bisa menyimpan:
    0.85
    atau
    85
    """
    value = _safe_number(value)

    if value > 1:
        value = value / 100

    return value


def _confidence_text(value):
    return f"{_confidence(value) * 100:.1f}%"


def _status(value):
    if value is None or value == "":
        return "Perlu dicek"

    text = str(value).strip()

    if text.lower() in [
        "success",
        "terbaca",
        "jelas",
        "valid",
        "berhasil",
        "detected",
    ]:
        return "Terbaca"

    if text.lower() in [
        "perlu dicek",
        "check",
        "need_check",
        "kurang jelas",
        "unclear",
        "gagal",
        "failed",
    ]:
        return "Perlu dicek"

    return text


def _format_datetime(value):
    if value is None or value == "":
        return "-"

    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M:%S")

    return str(value)


def _project_roots():
    """
    Folder-folder yang mungkin menjadi root project. URL browser
    '/static/captures/x.jpg' = <root>/static/captures/x.jpg di disk.
    Root Flask (current_app.root_path) dicoba lebih dulu, lalu folder
    file ini, induknya, dan working directory -- supaya tetap benar
    walau report_export.py dipindah ke subfolder.
    """
    roots = []

    try:
        from flask import current_app

        roots.append(current_app.root_path)
    except Exception:
        pass

    roots.append(BASE_DIR)
    roots.append(os.path.dirname(BASE_DIR))
    roots.append(os.getcwd())

    unique = []

    for root in roots:
        root = os.path.abspath(root)

        if root not in unique:
            unique.append(root)

    return unique


def _resolve_image_path(path):
    """
    Mencari file gambar dari path yang sama dengan yang dipakai website
    (website memuat foto lewat '/' + path, mis. '/static/captures/...').
    Mengembalikan path file lokal, atau None jika benar-benar tidak ada.
    """
    if not path:
        return None

    path = str(path).strip()

    if not path:
        return None

    # buang query string / fragment dan decode %20 dari URL browser
    path = unquote(path.split("?", 1)[0].split("#", 1)[0])
    path = path.replace("\\", "/")

    # path absolut di disk
    if os.path.isabs(path) and os.path.isfile(path):
        return path

    relative = path.lstrip("/")
    normalized = relative.replace("/", os.sep)

    subfolders = [
        "",
        "static",
        "captures",
        os.path.join("static", "captures"),
        os.path.join("static", "captures", "plates"),
        os.path.join("static", "captures", "faces"),
        os.path.join("static", "captures", "vehicles"),
    ]

    roots = _project_roots()

    for root in roots:
        for sub in subfolders:
            candidate = os.path.join(root, sub, normalized)

            if os.path.isfile(candidate):
                return candidate

    # jika yang tersimpan hanya nama file
    basename = os.path.basename(normalized)

    if basename:
        for root in roots:
            for sub in subfolders:
                if not sub:
                    continue

                candidate = os.path.join(root, sub, basename)

                if os.path.isfile(candidate):
                    return candidate

    return None


# ============================================================
# HELPER DATA WEBSITE  (PATCH: data export = data website)
#
# Semua nama field di bawah DIAMBIL dari kode render website
# (static/js/app.js), bukan tebakan:
#   loadDetections()   -> GET /api/detections
#   loadPlateHistory() -> GET /api/plate/history
#   loadRecap()        -> GET /api/analytics  (summary, cameras, daily)
#   loadEnterpriseStatistics() -> GET /api/analytics
#                                 (summary, cameras, hourly, top_plates)
# ============================================================

def _text(value, default="-"):
    """Teks apa adanya dari data. '-' hanya jika data memang kosong."""
    if value is None:
        return default

    text = str(value).strip()

    return text if text else default


def _num(value):
    """Angka float atau None (mengerti '74.6' dan '74.6%')."""
    if value is None or value == "" or isinstance(value, bool):
        return None

    try:
        return float(str(value).replace("%", "").strip())
    except (TypeError, ValueError):
        return None


def _int(value, default=0):
    number = _num(value)

    return int(round(number)) if number is not None else default


def _pct_text(percent):
    """Angka skala 0-100 -> '74.6%'. None -> '-'."""
    if percent is None:
        return "-"

    return f"{percent:.1f}%"


def _time_text(value):
    """
    Waktu deteksi ASLI dari data (website menampilkan item.timestamp
    apa adanya, contoh: 2026-09-21 09:49:19). Bukan waktu export.
    """
    if value is None or value == "":
        return "-"

    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M:%S")

    return str(value)


def _confidence_percent(item):
    """
    Confidence dalam skala 0-100 -- sama dengan item.confidence_percent
    yang dicetak website (`${item.confidence_percent}%`).

    Hanya jika field itu tidak ada, jatuh ke item.confidence dengan
    deteksi skala: 0-1 dikali 100, sisanya dianggap sudah persen.
    Tidak ada threshold / hitung ulang.
    """
    if not isinstance(item, dict):
        return None

    for key in (
        "confidence_percent",
        "vehicle_confidence_percent",
        "person_confidence_percent",
        "plate_confidence_percent",
    ):
        number = _num(item.get(key))

        if number is not None:
            return number

    raw = _num(item.get("confidence"))

    if raw is None:
        return None

    return raw * 100 if 0 <= raw <= 1 else raw


def _status_label(item):
    """
    Status persis seperti badge website: status_code 1 = Terbaca,
    2 = Perlu Cek, selain itu Gagal. Jika status_code tidak ada,
    pakai teks item.status apa adanya.
    """
    code = _num(item.get("status_code")) if isinstance(item, dict) else None

    if code is not None:
        code = int(code)

        if code == 1:
            return "Terbaca"

        if code == 2:
            return "Perlu Cek"

        return "Gagal"

    return _text(item.get("status") if isinstance(item, dict) else None)


def _detection_fields(item):
    """
    Salinan logika loadDetections() di app.js:

      isVehicle = object_type == 'vehicle' || type in
                  (vehicle, vehicle_with_plate, plate)
      isPlate   = type == 'vehicle_with_plate' || (isVehicle && has_plate)
      Target    = isVehicle ? (isPlate ? plate : 'Kendaraan') : 'Orang'
      Tipe      = isPlate ? 'Kendaraan / Plat Nomor'
                          : (isVehicle ? 'Kendaraan' : 'Orang')
      Foto      = isVehicle ? vehicle_image_path || plate_image_path
                            : face_image_path
    """
    object_type = item.get("object_type")
    type_ = item.get("type")

    is_vehicle = (
        object_type == "vehicle"
        or type_ in ("vehicle", "vehicle_with_plate", "plate")
    )

    is_plate = type_ == "vehicle_with_plate" or (
        is_vehicle and bool(item.get("has_plate"))
    )

    plate = _text(item.get("plate"), default="")

    if is_vehicle:
        target = plate if (is_plate and plate not in ("", "-")) else "Kendaraan"
    else:
        target = "Orang"

    if is_plate:
        kind = "Kendaraan / Plat Nomor"
    elif is_vehicle:
        kind = "Kendaraan"
    else:
        kind = "Orang"

    if is_vehicle:
        photo = (
            item.get("vehicle_image_path")
            or item.get("vehicleImagePath")
            or item.get("plate_image_path")
            or item.get("plateImagePath")
            or ""
        )
    else:
        photo = (
            item.get("face_image_path")
            or item.get("faceImagePath")
            or ""
        )

    confidence = _confidence_percent(item)

    return {
        "target": target,
        "kind": kind,
        "camera": _text(item.get("camera")),
        "confidence": confidence,
        "confidence_text": _pct_text(confidence),
        "time": _time_text(item.get("timestamp")),
        "status": _status_label(item),
        "photo": photo,
    }


def _plate_fields(item):
    """
    Salinan logika loadPlateHistory() di app.js:
    plate, camera, confidence_percent, timestamp, status_code,
    crop = item.image_path
    """
    confidence = _confidence_percent(item)

    return {
        "plate": _text(item.get("plate")),
        "camera": _text(item.get("camera")),
        "confidence": confidence,
        "confidence_text": _pct_text(confidence),
        "time": _time_text(item.get("timestamp")),
        "status": _status_label(item),
        "photo": item.get("image_path") or "",
    }


PERIOD_LABELS = {
    "today": "Hari Ini",
    "2d": "2 Hari Terakhir",
    "7d": "7 Hari Terakhir",
    "30d": "30 Hari Terakhir",
    "all": "Semua Waktu",
}


def _period_label(period):
    if not period:
        return ""

    return PERIOD_LABELS.get(str(period), str(period))


def _now_text():
    return datetime.now().strftime("%d-%m-%Y %H:%M:%S")


def _analytics_parts(analytics):
    """summary / cameras / daily / hourly / top_plates dari /api/analytics."""
    data = analytics if isinstance(analytics, dict) else {}

    return (
        data.get("summary") or {},
        data.get("cameras") or [],
        data.get("daily") or [],
        data.get("hourly") or [],
        data.get("top_plates") or [],
    )


def _peak_hour(hourly):
    """Sama dengan `peak` di loadEnterpriseStatistics(): jam dengan
    kendaraan terbanyak (yang pertama jika seri)."""
    best = None

    for item in hourly or []:
        vehicles = _num(item.get("vehicles")) or 0

        if best is None or vehicles > (_num(best.get("vehicles")) or 0):
            best = item

    return best


def _peak_label(item):
    label = _text(item.get("label"), default="")
    hour = _num(item.get("hour"))

    if hour is not None:
        return f"{label} - {(int(hour) + 1) % 24:02d}:00"

    return label or "-"


# ------------------------------------------------------------
# EXCEL - helper patch
# ------------------------------------------------------------

def _excel_table_sheet(
    ws,
    title,
    subtitle,
    headers,
    rows,
    widths,
    row_height=None,
    print_setup=False,
):
    """
    Pola sheet yang sama dengan export sebelumnya: title (baris 1),
    subtitle (baris 2), header (baris 4), data mulai baris 5,
    freeze panes, autofilter, lebar kolom, border.
    Data kosong -> header tetap dibuat.
    """
    columns = len(headers)

    _add_excel_title(ws, title, subtitle, columns)

    for col, value in enumerate(headers, start=1):
        ws.cell(row=4, column=col, value=value)

    _apply_excel_header(ws, 4, 1, columns)

    for index, row in enumerate(rows, start=1):
        row_number = 4 + index

        for col, value in enumerate(row, start=1):
            ws.cell(row=row_number, column=col, value=value)

        if row_height:
            ws.row_dimensions[row_number].height = row_height

    last_row = max(5, 4 + len(rows))

    _set_widths(ws, widths)

    _style_excel_sheet(ws)

    ws.auto_filter.ref = f"A4:{get_column_letter(columns)}{last_row}"

    if print_setup:
        ws.print_title_rows = "1:4"
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True

    return last_row


def _excel_percent_cell(cell, percent):
    """Confidence sebagai angka persen asli (0.746 -> tampil 74.6%)."""
    if percent is None:
        cell.value = "-"
        return

    cell.value = round(percent / 100, 6)
    cell.number_format = "0.0%"

    _confidence_excel_fill(cell, percent / 100)


def _excel_section_title(ws, row, text, columns):
    ws.merge_cells(
        start_row=row,
        start_column=1,
        end_row=row,
        end_column=columns,
    )

    cell = ws.cell(row=row, column=1, value=text)

    cell.font = Font(bold=True, size=12, color=WHITE)
    cell.fill = PatternFill("solid", fgColor=DARK)
    cell.alignment = Alignment(horizontal="left", vertical="center")

    ws.row_dimensions[row].height = 22


def _excel_block(ws, start_row, headers, rows):
    """Tulis header + baris data mulai start_row. Return baris kosong berikutnya."""
    columns = len(headers)

    for col, value in enumerate(headers, start=1):
        ws.cell(row=start_row, column=col, value=value)

    _apply_excel_header(ws, start_row, 1, columns)

    for offset, row in enumerate(rows, start=1):
        for col, value in enumerate(row, start=1):
            ws.cell(row=start_row + offset, column=col, value=value)

    return start_row + len(rows) + 1


# ------------------------------------------------------------
# PDF - helper patch
# ------------------------------------------------------------

EMPTY_MESSAGE = "Tidak ada data pada periode yang dipilih."


def _p(value, style):
    """Paragraph aman: karakter & < > pada nama kamera/plat tidak merusak PDF."""
    return Paragraph(escape(str(value)), style)


def _p_bold(value, style):
    return Paragraph(f"<b>{escape(str(value))}</b>", style)


def _pdf_styles():
    styles = getSampleStyleSheet()

    return {
        "title": ParagraphStyle(
            "PatchTitle",
            parent=styles["Title"],
            fontSize=17,
            leading=21,
            alignment=TA_CENTER,
            spaceAfter=3 * mm,
        ),
        "subtitle": ParagraphStyle(
            "PatchSubtitle",
            parent=styles["Normal"],
            fontSize=9,
            leading=12,
            alignment=TA_CENTER,
            spaceAfter=6 * mm,
        ),
        "section": ParagraphStyle(
            "PatchSection",
            parent=styles["Heading2"],
            fontSize=12,
            leading=15,
            spaceBefore=5 * mm,
            spaceAfter=3 * mm,
        ),
        "cell": ParagraphStyle(
            "PatchCell",
            parent=styles["Normal"],
            fontSize=8,
            leading=10,
            alignment=TA_LEFT,
        ),
        "head": ParagraphStyle(
            "PatchHead",
            parent=styles["Normal"],
            fontSize=8,
            leading=10,
            alignment=TA_CENTER,
            textColor=colors.white,
            fontName="Helvetica-Bold",
        ),
        "empty": ParagraphStyle(
            "PatchEmpty",
            parent=styles["Normal"],
            fontSize=10,
            leading=14,
            alignment=TA_CENTER,
            textColor=colors.HexColor("#6B7280"),
            spaceBefore=8 * mm,
        ),
    }


def _pdf_head(labels, style):
    return [_p(label, style) for label in labels]


def _pdf_table(data, col_widths, extra=None):
    """Tabel standar: header biru berulang tiap halaman, grid, zebra."""
    table = Table(data, colWidths=col_widths, repeatRows=1)

    style = [
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2563EB")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#D1D5DB")),
        (
            "ROWBACKGROUNDS",
            (0, 1),
            (-1, -1),
            [colors.white, colors.HexColor("#F9FAFB")],
        ),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ]

    style.extend(extra or [])

    table.setStyle(TableStyle(style))

    return table


def _pdf_footer(canvas, doc):
    canvas.saveState()
    canvas.setFont("Helvetica", 8)
    canvas.setFillColor(colors.HexColor("#6B7280"))
    canvas.drawRightString(
        doc.pagesize[0] - doc.rightMargin,
        8 * mm,
        f"PlateVision - Halaman {canvas.getPageNumber()}",
    )
    canvas.restoreState()


def _pdf_build(doc, story):
    doc.build(story, onFirstPage=_pdf_footer, onLaterPages=_pdf_footer)


def _pdf_photo(path, cell_style, max_w, max_h, missing_text):
    """
    Gambar capture asli untuk sel tabel PDF.

    Path dari database/website (mis. 'static/captures/...') diubah lebih
    dulu menjadi path FILE LOKAL oleh _resolve_image_path(), karena
    reportlab tidak bisa membuka URL browser. Gambar diperkecil supaya
    PDF tidak membengkak. Teks 'tidak tersedia' hanya muncul jika file
    memang tidak ditemukan / tidak bisa dibaca.
    """
    resolved = _resolve_image_path(path)

    if not resolved:
        return _p(missing_text, cell_style)

    try:
        source = resolved

        try:
            from PIL import Image as PILImage

            with PILImage.open(resolved) as img:
                img.load()
                width, height = img.size

                if img.mode not in ("RGB", "L"):
                    img = img.convert("RGB")

                # cukup ~ 300 dpi untuk sel kecil
                limit = 500
                if max(width, height) > limit:
                    img = img.copy()
                    img.thumbnail((limit, limit))

                data = io.BytesIO()
                img.save(data, format="JPEG", quality=85)
                data.seek(0)
                source = data
        except ImportError:
            reader = ImageReader(resolved)
            width, height = reader.getSize()

        ratio = min(max_w / width, max_h / height)

        return Image(source, width=width * ratio, height=height * ratio)
    except Exception:
        return _p(missing_text, cell_style)


# ============================================================
# EXCEL HELPER
# ============================================================

def _apply_excel_header(ws, row, start_col, end_col):
    fill = PatternFill("solid", fgColor=HEADER)

    for col in range(start_col, end_col + 1):
        cell = ws.cell(row=row, column=col)

        cell.fill = fill
        cell.font = Font(
            color=WHITE,
            bold=True,
            size=11,
        )
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )

        cell.border = Border(
            bottom=Side(
                style="thin",
                color=BORDER_COLOR,
            )
        )


def _style_excel_sheet(ws):
    ws.freeze_panes = "A4"
    ws.sheet_view.showGridLines = False

    for row in ws.iter_rows():
        for cell in row:
            # title / header / judul seksi (bold atau italic) tetap
            # memakai font yang sudah diberi _add_excel_title / _apply_excel_header
            if not (cell.font and (cell.font.bold or cell.font.italic)):
                cell.font = Font(
                    name="Calibri",
                    size=10,
                    color=TEXT,
                )

            cell.alignment = Alignment(
                horizontal=cell.alignment.horizontal,
                vertical="center",
                wrap_text=True,
            )

            cell.border = Border(
                left=Side(style="thin", color=BORDER_COLOR),
                right=Side(style="thin", color=BORDER_COLOR),
                top=Side(style="thin", color=BORDER_COLOR),
                bottom=Side(style="thin", color=BORDER_COLOR),
            )


def _add_excel_title(ws, title, subtitle, columns):
    ws.merge_cells(
        start_row=1,
        start_column=1,
        end_row=1,
        end_column=columns,
    )

    title_cell = ws.cell(row=1, column=1)
    title_cell.value = title
    title_cell.font = Font(
        size=16,
        bold=True,
        color=WHITE,
    )
    title_cell.fill = PatternFill(
        "solid",
        fgColor=DARK,
    )
    title_cell.alignment = Alignment(
        horizontal="center",
        vertical="center",
    )

    ws.row_dimensions[1].height = 28

    ws.merge_cells(
        start_row=2,
        start_column=1,
        end_row=2,
        end_column=columns,
    )

    subtitle_cell = ws.cell(row=2, column=1)
    subtitle_cell.value = subtitle
    subtitle_cell.font = Font(
        size=10,
        color=TEXT,
        italic=True,
    )
    subtitle_cell.fill = PatternFill(
        "solid",
        fgColor=LIGHT_GRAY,
    )
    subtitle_cell.alignment = Alignment(
        horizontal="center",
        vertical="center",
    )

    ws.row_dimensions[2].height = 22


def _set_widths(ws, widths):
    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[
            get_column_letter(index)
        ].width = width


def _status_excel_fill(cell, status):
    text = str(status).lower()

    if "terbaca" in text or "jelas" in text or "berhasil" in text:
        cell.fill = PatternFill(
            "solid",
            fgColor=LIGHT_GREEN,
        )

    elif "dicek" in text or "unclear" in text or "gagal" in text:
        cell.fill = PatternFill(
            "solid",
            fgColor=LIGHT_YELLOW,
        )


def _confidence_excel_fill(cell, confidence):
    value = _confidence(confidence)

    if value >= 0.70:
        cell.fill = PatternFill(
            "solid",
            fgColor=LIGHT_GREEN,
        )
    else:
        cell.fill = PatternFill(
            "solid",
            fgColor=LIGHT_YELLOW,
        )


def _add_excel_table(ws, ref, name):
    try:
        table = Table(
            displayName=name,
            ref=ref,
        )

        style = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )

        table.tableStyleInfo = style
        ws.add_table(table)
    except Exception:
        pass


def _finalize_excel(wb):
    buffer = io.BytesIO()

    wb.save(buffer)

    buffer.seek(0)

    return buffer


# ============================================================
# 1. HASIL DETEKSI - EXCEL
#    Kolom: No | Target / Nilai | Jenis Deteksi | Area CCTV |
#           Confidence | Waktu Deteksi | Status
#    Tanpa Keterangan, tanpa Foto, tanpa Aksi.
#    Sumber field: lihat _detection_fields() (= loadDetections()).
# ============================================================

def build_detections_excel(items):
    wb = Workbook()

    ws = wb.active
    ws.title = "Hasil Deteksi"

    headers = [
        "No",
        "Target / Nilai",
        "Jenis Deteksi",
        "Area CCTV",
        "Confidence",
        "Waktu Deteksi",
        "Status",
    ]

    fields = [_detection_fields(item) for item in (items or [])]

    rows = [
        [
            index,
            f["target"],
            f["kind"],
            f["camera"],
            None,  # confidence: diisi sebagai persen numerik di bawah
            f["time"],
            f["status"],
        ]
        for index, f in enumerate(fields, start=1)
    ]

    _excel_table_sheet(
        ws,
        "HASIL DETEKSI CCTV",
        "Laporan hasil deteksi wajah dan plat nomor kendaraan "
        f"— Diekspor: {_now_text()} — Jumlah data: {len(rows)}",
        headers,
        rows,
        [8, 24, 26, 28, 15, 22, 16],
        row_height=24,
        print_setup=True,
    )

    for index, f in enumerate(fields, start=1):
        row_number = 4 + index

        _excel_percent_cell(ws.cell(row=row_number, column=5), f["confidence"])

        _status_excel_fill(ws.cell(row=row_number, column=7), f["status"])

    return _finalize_excel(wb)


# ============================================================
# 2. HASIL DETEKSI - PDF
#    Kolom: No | Foto | Target / Nilai | Jenis Deteksi | Area CCTV |
#           Confidence | Waktu Deteksi | Status   (tanpa Keterangan)
# ============================================================

def build_detections_pdf(items):
    buffer = io.BytesIO()

    doc = SimpleDocTemplate(
        buffer,
        pagesize=landscape(A4),
        rightMargin=12 * mm,
        leftMargin=12 * mm,
        topMargin=12 * mm,
        bottomMargin=14 * mm,
    )

    s = _pdf_styles()

    fields = [_detection_fields(item) for item in (items or [])]

    story = [
        Paragraph("HASIL DETEKSI CCTV", s["title"]),
        Paragraph(
            "Laporan hasil deteksi wajah dan plat nomor kendaraan"
            f"<br/>Dicetak: {_now_text()} — Jumlah data: {len(fields)}",
            s["subtitle"],
        ),
    ]

    if not fields:
        story.append(Paragraph(EMPTY_MESSAGE, s["empty"]))
        _pdf_build(doc, story)
        buffer.seek(0)
        return buffer

    data = [
        _pdf_head(
            [
                "No",
                "Foto",
                "Target / Nilai",
                "Jenis Deteksi",
                "Area CCTV",
                "Confidence",
                "Waktu Deteksi",
                "Status",
            ],
            s["head"],
        )
    ]

    extra = [
        ("ALIGN", (0, 1), (0, -1), "CENTER"),
        ("ALIGN", (1, 1), (1, -1), "CENTER"),
    ]

    for index, f in enumerate(fields, start=1):
        data.append(
            [
                str(index),
                _pdf_photo(
                    f["photo"],
                    s["cell"],
                    30 * mm,
                    18 * mm,
                    "Foto tidak tersedia",
                ),
                _p(f["target"], s["cell"]),
                _p(f["kind"], s["cell"]),
                _p(f["camera"], s["cell"]),
                _p(f["confidence_text"], s["cell"]),
                _p(f["time"], s["cell"]),
                _p(f["status"], s["cell"]),
            ]
        )

        row = len(data) - 1

        confident = (f["confidence"] or 0) >= 70

        extra.append(
            (
                "BACKGROUND",
                (5, row),
                (5, row),
                colors.HexColor("#DCFCE7" if confident else "#FEF3C7"),
            )
        )

        extra.append(
            (
                "BACKGROUND",
                (7, row),
                (7, row),
                colors.HexColor(
                    "#DCFCE7" if f["status"] == "Terbaca" else "#FEF3C7"
                ),
            )
        )

    story.append(
        _pdf_table(
            data,
            [
                10 * mm,
                40 * mm,
                38 * mm,
                40 * mm,
                44 * mm,
                26 * mm,
                40 * mm,
                28 * mm,
            ],
            extra,
        )
    )

    _pdf_build(doc, story)

    buffer.seek(0)

    return buffer


# ============================================================
# 3. RIWAYAT PLAT - EXCEL
#    Kolom: No | Nomor Plat | Area CCTV | Confidence |
#           Waktu Deteksi | Status   (tanpa Keterangan, tanpa crop)
#    Sumber field: lihat _plate_fields() (= loadPlateHistory()).
# ============================================================

def build_plates_excel(items):
    wb = Workbook()

    ws = wb.active
    ws.title = "Riwayat Plat"

    headers = [
        "No",
        "Nomor Plat",
        "Area CCTV",
        "Confidence",
        "Waktu Deteksi",
        "Status",
    ]

    fields = [_plate_fields(item) for item in (items or [])]

    rows = [
        [
            index,
            f["plate"],
            f["camera"],
            None,
            f["time"],
            f["status"],
        ]
        for index, f in enumerate(fields, start=1)
    ]

    _excel_table_sheet(
        ws,
        "RIWAYAT PLAT NOMOR",
        "Riwayat hasil pembacaan plat kendaraan dari CCTV "
        f"— Diekspor: {_now_text()} — Jumlah data: {len(rows)}",
        headers,
        rows,
        [8, 22, 32, 16, 23, 16],
        row_height=24,
        print_setup=True,
    )

    for index, f in enumerate(fields, start=1):
        row_number = 4 + index

        _excel_percent_cell(ws.cell(row=row_number, column=4), f["confidence"])

        _status_excel_fill(ws.cell(row=row_number, column=6), f["status"])

    return _finalize_excel(wb)


# ============================================================
# 4. RIWAYAT PLAT - PDF
#    Kolom: No | Crop / Foto Plat | Nomor Plat | Area CCTV |
#           Confidence | Waktu Deteksi | Status  (tanpa Keterangan)
# ============================================================

def build_plates_pdf(items):
    buffer = io.BytesIO()

    doc = SimpleDocTemplate(
        buffer,
        pagesize=landscape(A4),
        rightMargin=12 * mm,
        leftMargin=12 * mm,
        topMargin=12 * mm,
        bottomMargin=14 * mm,
    )

    s = _pdf_styles()

    fields = [_plate_fields(item) for item in (items or [])]

    story = [
        Paragraph("RIWAYAT PLAT NOMOR", s["title"]),
        Paragraph(
            "Riwayat hasil pembacaan plat kendaraan dari CCTV"
            f"<br/>Dicetak: {_now_text()} — Jumlah data: {len(fields)}",
            s["subtitle"],
        ),
    ]

    if not fields:
        story.append(Paragraph(EMPTY_MESSAGE, s["empty"]))
        _pdf_build(doc, story)
        buffer.seek(0)
        return buffer

    data = [
        _pdf_head(
            [
                "No",
                "Crop / Foto Plat",
                "Nomor Plat",
                "Area CCTV",
                "Confidence",
                "Waktu Deteksi",
                "Status",
            ],
            s["head"],
        )
    ]

    extra = [
        ("ALIGN", (0, 1), (0, -1), "CENTER"),
        ("ALIGN", (1, 1), (1, -1), "CENTER"),
    ]

    for index, f in enumerate(fields, start=1):
        data.append(
            [
                str(index),
                _pdf_photo(
                    f["photo"],
                    s["cell"],
                    38 * mm,
                    16 * mm,
                    "Crop tidak tersedia",
                ),
                _p_bold(f["plate"], s["cell"]),
                _p(f["camera"], s["cell"]),
                _p(f["confidence_text"], s["cell"]),
                _p(f["time"], s["cell"]),
                _p(f["status"], s["cell"]),
            ]
        )

        row = len(data) - 1

        confident = (f["confidence"] or 0) >= 70

        extra.append(
            (
                "BACKGROUND",
                (4, row),
                (4, row),
                colors.HexColor("#DCFCE7" if confident else "#FEF3C7"),
            )
        )

        extra.append(
            (
                "BACKGROUND",
                (6, row),
                (6, row),
                colors.HexColor(
                    "#DCFCE7" if f["status"] == "Terbaca" else "#FEF3C7"
                ),
            )
        )

    story.append(
        _pdf_table(
            data,
            [
                10 * mm,
                48 * mm,
                42 * mm,
                50 * mm,
                30 * mm,
                48 * mm,
                34 * mm,
            ],
            extra,
        )
    )

    _pdf_build(doc, story)

    buffer.seek(0)

    return buffer


# ============================================================
# 5. REKAPITULASI - EXCEL   (1 file, 3 sheet)
#    Sheet 1 Ringkasan    : Kendaraan/Orang Masuk, Keluar, Total
#    Sheet 2 Rekap CCTV   : CCTV / Gerbang | Arah | Kendaraan |
#                           Orang | Plat Unik | Status
#    Sheet 3 Total Harian : Tanggal | Kendaraan | Plat Unik |
#                           Masuk | Keluar | Orang
#    Sumber field: loadRecap() -> /api/analytics
# ============================================================

def _recap_summary_rows(summary):
    return [
        ["Kendaraan Masuk", _int(summary.get("vehicle_entry"))],
        ["Kendaraan Keluar", _int(summary.get("vehicle_exit"))],
        ["Total Kendaraan", _int(summary.get("vehicles"))],
        ["Orang Masuk", _int(summary.get("people_entry"))],
        ["Orang Keluar", _int(summary.get("people_exit"))],
        ["Total Orang", _int(summary.get("people"))],
    ]


def _recap_camera_rows(cameras):
    return [
        [
            _text(c.get("camera")),
            _text(c.get("direction")),
            _int(c.get("vehicles")),
            _int(c.get("people")),
            _int(c.get("unique_plates")),
            _text(c.get("status")),
        ]
        for c in cameras
    ]


def _recap_daily_rows(daily):
    return [
        [
            _text(d.get("date")),
            _int(d.get("vehicles")),
            _int(d.get("unique_plates")),
            _int(d.get("entry")),
            _int(d.get("exit")),
            _int(d.get("people")),
        ]
        for d in daily
    ]


def _recap_subtitle(analytics, period, text):
    label = _period_label(
        period
        or (analytics.get("period") if isinstance(analytics, dict) else "")
    )

    parts = [text]

    if label:
        parts.append(f"Periode: {label}")

    parts.append(f"Diekspor: {_now_text()}")

    return " — ".join(parts)


def build_recap_excel(analytics, period=None):
    wb = Workbook()

    summary, cameras, daily, _hourly, _top = _analytics_parts(analytics)

    # --------------------------------------------------------
    # SHEET 1 - RINGKASAN
    # --------------------------------------------------------

    ws = wb.active
    ws.title = "Ringkasan"

    _excel_table_sheet(
        ws,
        "REKAPITULASI MONITORING CCTV",
        _recap_subtitle(
            analytics, period, "Ringkasan aktivitas kendaraan dan orang"
        ),
        ["Indikator", "Jumlah"],
        _recap_summary_rows(summary),
        [30, 18],
    )

    for row in range(5, 5 + 6):
        ws.cell(row=row, column=2).number_format = "#,##0"
        ws.cell(row=row, column=2).alignment = Alignment(horizontal="center")

    ws.auto_filter.ref = None

    # --------------------------------------------------------
    # SHEET 2 - REKAP CCTV
    # --------------------------------------------------------

    ws_camera = wb.create_sheet("Rekap CCTV")

    camera_rows = _recap_camera_rows(cameras)

    _excel_table_sheet(
        ws_camera,
        "REKAP PER CCTV / GERBANG",
        _recap_subtitle(
            analytics,
            period,
            "Agregasi event deteksi per CCTV / gerbang",
        ),
        [
            "CCTV / Gerbang",
            "Arah",
            "Kendaraan",
            "Orang",
            "Plat Unik",
            "Status",
        ],
        camera_rows,
        [34, 20, 14, 14, 14, 14],
    )

    for index, row in enumerate(camera_rows, start=1):
        cell = ws_camera.cell(row=4 + index, column=6)

        cell.fill = PatternFill(
            "solid",
            fgColor=LIGHT_GREEN if row[5] == "Aktif" else LIGHT_YELLOW,
        )

    # --------------------------------------------------------
    # SHEET 3 - TOTAL HARIAN
    # --------------------------------------------------------

    ws_daily = wb.create_sheet("Total Harian")

    _excel_table_sheet(
        ws_daily,
        "TOTAL PER HARI",
        _recap_subtitle(
            analytics,
            period,
            "Kendaraan, plat unik, masuk/keluar, dan orang per tanggal",
        ),
        [
            "Tanggal",
            "Kendaraan",
            "Plat Unik",
            "Masuk",
            "Keluar",
            "Orang",
        ],
        _recap_daily_rows(daily),
        [16, 14, 14, 14, 14, 14],
    )

    return _finalize_excel(wb)


# ============================================================
# 6. REKAPITULASI - PDF
#    Sumber data identik dengan build_recap_excel().
# ============================================================

def build_recap_pdf(analytics, period=None):
    buffer = io.BytesIO()

    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        rightMargin=15 * mm,
        leftMargin=15 * mm,
        topMargin=15 * mm,
        bottomMargin=15 * mm,
    )

    s = _pdf_styles()

    summary, cameras, daily, _hourly, _top = _analytics_parts(analytics)

    label = _period_label(
        period
        or (analytics.get("period") if isinstance(analytics, dict) else "")
    )

    subtitle = "Ringkasan hasil pemantauan dan deteksi CCTV"

    if label:
        subtitle += f"<br/>Periode: {escape(label)}"

    subtitle += f"<br/>Dicetak: {_now_text()}"

    story = [
        Paragraph("REKAPITULASI MONITORING CCTV", s["title"]),
        Paragraph(subtitle, s["subtitle"]),
    ]

    # ---- 1. Ringkasan ----
    story.append(Paragraph("Ringkasan", s["section"]))

    summary_data = [_pdf_head(["Indikator", "Jumlah"], s["head"])]

    for name, value in _recap_summary_rows(summary):
        summary_data.append([_p(name, s["cell"]), str(value)])

    story.append(
        _pdf_table(
            summary_data,
            [100 * mm, 50 * mm],
            [("ALIGN", (1, 1), (1, -1), "CENTER")],
        )
    )

    # ---- 2. Rekap per CCTV / Gerbang ----
    story.append(Paragraph("Rekap per CCTV / Gerbang", s["section"]))

    camera_rows = _recap_camera_rows(cameras)

    if not camera_rows:
        story.append(Paragraph(EMPTY_MESSAGE, s["empty"]))
    else:
        camera_data = [
            _pdf_head(
                [
                    "CCTV / Gerbang",
                    "Arah",
                    "Kendaraan",
                    "Orang",
                    "Plat Unik",
                    "Status",
                ],
                s["head"],
            )
        ]

        extra = [("ALIGN", (2, 1), (5, -1), "CENTER")]

        for row in camera_rows:
            camera_data.append(
                [
                    _p(row[0], s["cell"]),
                    _p(row[1], s["cell"]),
                    str(row[2]),
                    str(row[3]),
                    str(row[4]),
                    _p(row[5], s["cell"]),
                ]
            )

            index = len(camera_data) - 1

            extra.append(
                (
                    "BACKGROUND",
                    (5, index),
                    (5, index),
                    colors.HexColor(
                        "#DCFCE7" if row[5] == "Aktif" else "#FEF3C7"
                    ),
                )
            )

        story.append(
            _pdf_table(
                camera_data,
                [
                    50 * mm,
                    32 * mm,
                    24 * mm,
                    20 * mm,
                    24 * mm,
                    26 * mm,
                ],
                extra,
            )
        )

    # ---- 3. Total per Hari ----
    story.append(Paragraph("Total per Hari", s["section"]))

    daily_rows = _recap_daily_rows(daily)

    if not daily_rows:
        story.append(Paragraph(EMPTY_MESSAGE, s["empty"]))
    else:
        daily_data = [
            _pdf_head(
                [
                    "Tanggal",
                    "Kendaraan",
                    "Plat Unik",
                    "Masuk",
                    "Keluar",
                    "Orang",
                ],
                s["head"],
            )
        ]

        for row in daily_rows:
            daily_data.append(
                [_p(row[0], s["cell"])] + [str(value) for value in row[1:]]
            )

        story.append(
            _pdf_table(
                daily_data,
                [34 * mm, 28 * mm, 28 * mm, 28 * mm, 28 * mm, 28 * mm],
                [("ALIGN", (1, 1), (-1, -1), "CENTER")],
            )
        )

    _pdf_build(doc, story)

    buffer.seek(0)

    return buffer


# ============================================================
# 7. STATISTIK - EXCEL   (1 file, 2 sheet: Statistik & Top 10)
#    Sumber field: loadEnterpriseStatistics() -> /api/analytics
#    (summary, cameras, hourly, top_plates). TIDAK memakai
#    struktur / fungsi Rekapitulasi.
# ============================================================

def _stats_summary_rows(summary):
    """Kartu KPI halaman Statistik."""
    vehicles = _num(summary.get("vehicles")) or 0
    plates = _num(summary.get("plates")) or 0

    read_rate = int(plates / vehicles * 100 + 0.5) if vehicles else 0

    return [
        ["Total Deteksi (Kendaraan)", _int(summary.get("vehicles"))],
        ["Deteksi Plat", _int(summary.get("plates"))],
        ["Read Rate Plat", f"{read_rate}%"],
        ["Wajah / Orang", _int(summary.get("people"))],
        [
            "Ketersediaan CCTV (aktif / total)",
            f"{_int(summary.get('active_cameras'))} / "
            f"{_int(summary.get('total_cameras'))}",
        ],
    ]


def _stats_camera_rows(cameras):
    return [
        [
            _text(c.get("camera")),
            _int(c.get("vehicles")),
            _int(c.get("people")),
            _int(c.get("unique_plates")),
        ]
        for c in cameras
    ]


def _stats_hourly_rows(hourly):
    return [
        [
            _text(h.get("label")),
            _int(h.get("vehicles")),
            _int(h.get("people")),
        ]
        for h in hourly
    ]


def _stats_top_rows(top_plates):
    """Top 10. Mengenali dua bentuk data yang sama-sama dipakai website."""
    rows = []

    for plate in (top_plates or [])[:10]:
        confidence = _num(
            plate.get("confidence")
            if plate.get("confidence") is not None
            else plate.get("avg_confidence_percent")
        )

        rows.append(
            [
                _text(plate.get("plate") or plate.get("plate_number")),
                _int(
                    plate.get("count")
                    if plate.get("count") is not None
                    else plate.get("total_seen")
                ),
                _text(plate.get("camera") or plate.get("last_camera")),
                _time_text(plate.get("last_seen")),
                confidence,  # skala persen, sama seperti `${p.confidence}%`
                _status_label(plate),  # status_code/status record terakhir
            ]
        )

    return rows


def _stats_subtitle(period):
    label = _period_label(period)

    return f"Periode: {label}" if label else "Periode statistik"


def build_statistics_excel(analytics, period="today"):
    wb = Workbook()

    summary, cameras, _daily, hourly, top_plates = _analytics_parts(analytics)

    # --------------------------------------------------------
    # SHEET STATISTIK
    # --------------------------------------------------------

    ws = wb.active
    ws.title = "Statistik"

    _add_excel_title(
        ws,
        "STATISTIK MONITORING CCTV",
        f"{_stats_subtitle(period)} — Diekspor: {_now_text()}",
        4,
    )

    row = 4

    # Ringkasan Statistik
    _excel_section_title(ws, row, "RINGKASAN STATISTIK", 4)
    row = _excel_block(
        ws, row + 1, ["Indikator", "Nilai"], _stats_summary_rows(summary)
    )

    # Beban Lalu Lintas per CCTV
    row += 1
    _excel_section_title(ws, row, "BEBAN LALU LINTAS PER TITIK CCTV", 4)

    camera_rows = _stats_camera_rows(cameras)

    row = _excel_block(
        ws,
        row + 1,
        ["CCTV", "Kendaraan", "Orang", "Plat Unik"],
        camera_rows,
    )

    # Analisis Jam Sibuk
    row += 1
    _excel_section_title(ws, row, "ANALISIS JAM SIBUK (PEAK HOURS)", 4)

    peak = _peak_hour(hourly)

    row += 1

    if peak:
        ws.cell(row=row, column=1, value="Jam puncak")
        ws.cell(row=row, column=2, value=_peak_label(peak))
        ws.cell(
            row=row,
            column=3,
            value=(
                f"{_int(peak.get('vehicles'))} kendaraan · "
                f"{_int(peak.get('people'))} orang"
            ),
        )
        ws.merge_cells(start_row=row, start_column=3, end_row=row, end_column=4)
    else:
        ws.cell(row=row, column=1, value="Belum ada data.")

    row += 2

    row = _excel_block(
        ws,
        row,
        ["Jam", "Kendaraan", "Orang"],
        _stats_hourly_rows(hourly),
    )

    _set_widths(ws, [36, 22, 22, 18])

    _style_excel_sheet(ws)

    # --------------------------------------------------------
    # SHEET TOP 10
    # --------------------------------------------------------

    ws_top = wb.create_sheet("Top 10")

    top_rows = _stats_top_rows(top_plates)

    rows = [
        [index, plate, count, camera, seen, None, status]
        for index, (plate, count, camera, seen, _conf, status) in enumerate(
            top_rows, start=1
        )
    ]

    _excel_table_sheet(
        ws_top,
        "TOP 10 KENDARAAN / PLAT PALING SERING TERDETEKSI",
        f"{_stats_subtitle(period)} — Diekspor: {_now_text()}",
        [
            "Ranking",
            "Nomor Plat",
            "Jumlah Terdeteksi",
            "CCTV",
            "Terakhir Terlihat",
            "Confidence",
            "Status Terakhir",
        ],
        rows,
        [10, 22, 20, 30, 24, 14, 18],
    )

    for index, (_p1, _c, _cam, _seen, confidence, status) in enumerate(
        top_rows, start=1
    ):
        _excel_percent_cell(ws_top.cell(row=4 + index, column=6), confidence)

        if status == "Terbaca":
            ws_top.cell(row=4 + index, column=7).fill = PatternFill(
                "solid", fgColor=LIGHT_GREEN
            )
        elif status in ("Perlu Cek", "Gagal"):
            ws_top.cell(row=4 + index, column=7).fill = PatternFill(
                "solid", fgColor=LIGHT_YELLOW
            )

    return _finalize_excel(wb)


# ============================================================
# 8. STATISTIK - PDF   (backend reportlab + send_file, tanpa print)
#    Urutan: Judul, Periode, Tanggal cetak, Ringkasan Statistik,
#            Beban Lalu Lintas per CCTV, Analisis Jam Sibuk, Top 10
# ============================================================

def build_statistics_pdf(analytics, period="today"):
    buffer = io.BytesIO()

    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        rightMargin=15 * mm,
        leftMargin=15 * mm,
        topMargin=15 * mm,
        bottomMargin=15 * mm,
    )

    s = _pdf_styles()

    summary, cameras, _daily, hourly, top_plates = _analytics_parts(analytics)

    story = [
        Paragraph("STATISTIK MONITORING CCTV", s["title"]),
        Paragraph(
            f"{escape(_stats_subtitle(period))}"
            f"<br/>Tanggal cetak: {_now_text()}",
            s["subtitle"],
        ),
    ]

    # ---- 1. Ringkasan Statistik ----
    story.append(Paragraph("Ringkasan Statistik", s["section"]))

    summary_data = [_pdf_head(["Indikator", "Nilai"], s["head"])]

    for name, value in _stats_summary_rows(summary):
        summary_data.append([_p(name, s["cell"]), str(value)])

    story.append(
        _pdf_table(
            summary_data,
            [105 * mm, 50 * mm],
            [("ALIGN", (1, 1), (1, -1), "CENTER")],
        )
    )

    # ---- 2. Beban Lalu Lintas per CCTV ----
    story.append(Paragraph("Beban Lalu Lintas per Titik CCTV", s["section"]))

    camera_rows = _stats_camera_rows(cameras)

    if not camera_rows:
        story.append(Paragraph(EMPTY_MESSAGE, s["empty"]))
    else:
        camera_data = [
            _pdf_head(
                ["CCTV", "Kendaraan", "Orang", "Plat Unik"], s["head"]
            )
        ]

        for row in camera_rows:
            camera_data.append(
                [_p(row[0], s["cell"])] + [str(value) for value in row[1:]]
            )

        story.append(
            _pdf_table(
                camera_data,
                [75 * mm, 35 * mm, 35 * mm, 35 * mm],
                [("ALIGN", (1, 1), (-1, -1), "CENTER")],
            )
        )

    # ---- 3. Analisis Jam Sibuk ----
    story.append(Paragraph("Analisis Jam Sibuk (Peak Hours)", s["section"]))

    peak = _peak_hour(hourly)

    if not peak:
        story.append(Paragraph(EMPTY_MESSAGE, s["empty"]))
    else:
        story.append(
            Paragraph(
                f"<b>Jam puncak: {escape(_peak_label(peak))}</b> — "
                f"{_int(peak.get('vehicles'))} kendaraan · "
                f"{_int(peak.get('people'))} orang",
                s["cell"],
            )
        )

        story.append(Spacer(1, 3 * mm))

        hourly_data = [
            _pdf_head(["Jam", "Kendaraan", "Orang"], s["head"])
        ]

        extra = [("ALIGN", (0, 1), (-1, -1), "CENTER")]

        for row in _stats_hourly_rows(hourly):
            hourly_data.append([str(value) for value in row])

            if _text(peak.get("label")) == row[0]:
                extra.append(
                    (
                        "BACKGROUND",
                        (0, len(hourly_data) - 1),
                        (-1, len(hourly_data) - 1),
                        colors.HexColor("#FEF3C7"),
                    )
                )

        story.append(
            _pdf_table(
                hourly_data,
                [50 * mm, 50 * mm, 50 * mm],
                extra,
            )
        )

    # ---- 4. Top 10 (paling bawah) ----
    story.append(
        Paragraph("Top 10 Kendaraan / Plat Paling Sering Terdeteksi", s["section"])
    )

    top_rows = _stats_top_rows(top_plates)

    if not top_rows:
        story.append(Paragraph(EMPTY_MESSAGE, s["empty"]))
    else:
        top_data = [
            _pdf_head(
                [
                    "Ranking",
                    "Nomor Plat",
                    "Jumlah",
                    "CCTV",
                    "Terakhir Terlihat",
                    "Confidence",
                    "Status Terakhir",
                ],
                s["head"],
            )
        ]

        top_extra = [
            ("ALIGN", (0, 1), (0, -1), "CENTER"),
            ("ALIGN", (2, 1), (2, -1), "CENTER"),
            ("ALIGN", (5, 1), (5, -1), "CENTER"),
        ]

        for index, (plate, count, camera, seen, confidence, status) in enumerate(
            top_rows, start=1
        ):
            top_data.append(
                [
                    str(index),
                    _p_bold(plate, s["cell"]),
                    f"{count} kali",
                    _p(camera, s["cell"]),
                    _p(seen, s["cell"]),
                    _pct_text(confidence),
                    _p(status, s["cell"]),
                ]
            )

            if status == "Terbaca":
                top_extra.append(
                    ("BACKGROUND", (6, index), (6, index), colors.HexColor("#DCFCE7"))
                )
            elif status in ("Perlu Cek", "Gagal"):
                top_extra.append(
                    ("BACKGROUND", (6, index), (6, index), colors.HexColor("#FEF3C7"))
                )

        story.append(
            _pdf_table(
                top_data,
                [
                    16 * mm,
                    28 * mm,
                    20 * mm,
                    29 * mm,
                    36 * mm,
                    23 * mm,
                    28 * mm,
                ],
                top_extra,
            )
        )

    _pdf_build(doc, story)

    buffer.seek(0)

    return buffer


# ============================================================
# 9. MONITORING CCTV - EXCEL
#    Data berasal dari db.get_all_cameras() -- fungsi yang sama
#    dipakai endpoint GET /cameras -- supaya data website, Excel,
#    dan PDF selalu konsisten. Tidak ada foto/preview kamera.
# ============================================================

def _camera_status(value):
    """Ambil label status kamera dari data backend, jangan bikin status baru."""
    if value is None or value == "":
        return "Nonaktif"

    if isinstance(value, str):
        text = value.strip()
        return text if text else "Nonaktif"

    return "Aktif" if value else "Nonaktif"


def _camera_status_excel_fill(cell, status):
    text = str(status).lower()

    if "non" in text:
        cell.fill = PatternFill("solid", fgColor=LIGHT_GRAY)
    elif "error" in text:
        cell.fill = PatternFill("solid", fgColor=LIGHT_RED)
    elif "aktif" in text:
        cell.fill = PatternFill("solid", fgColor=LIGHT_GREEN)
    else:
        cell.fill = PatternFill("solid", fgColor=LIGHT_GRAY)


def build_cameras_excel(cameras_db):
    wb = Workbook()

    ws = wb.active
    ws.title = "Monitoring CCTV"

    columns = 4

    _add_excel_title(
        ws,
        "LAPORAN MONITORING CCTV",
        f"Diekspor pada {datetime.now().strftime('%d-%m-%Y %H:%M:%S')} "
        f"— Jumlah kamera: {len(cameras_db or [])}",
        columns,
    )

    headers = [
        "No",
        "Nama CCTV",
        "URL Stream",
        "Status",
    ]

    for col, value in enumerate(headers, start=1):
        ws.cell(row=4, column=col, value=value)

    _apply_excel_header(ws, 4, 1, columns)

    for index, camera in enumerate(cameras_db or [], start=1):
        name = _value(
            camera,
            "location",
            "name",
            "camera_name",
            default="-",
        )

        url = _value(
            camera,
            "stream_url",
            "url",
            default="-",
        )

        status = _camera_status(
            _value(
                camera,
                "status",
                "active",
                default=None,
            )
        )

        values = [
            index,
            _safe_text(name),
            _safe_text(url),
            status,
        ]

        row_number = 4 + index

        for col, value in enumerate(values, start=1):
            ws.cell(row=row_number, column=col, value=value)

        _camera_status_excel_fill(
            ws.cell(row=row_number, column=4),
            status,
        )

        ws.row_dimensions[row_number].height = 22

    last_row = max(5, 4 + len(cameras_db or []))

    _add_excel_table(
        ws,
        f"A4:D{last_row}",
        "CamerasTable",
    )

    _set_widths(ws, [8, 30, 60, 16])

    _style_excel_sheet(ws)

    ws.auto_filter.ref = f"A4:D{last_row}"

    return _finalize_excel(wb)


# ============================================================
# 10. MONITORING CCTV - PDF
# ============================================================

def build_cameras_pdf(cameras_db):
    buffer = io.BytesIO()

    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        rightMargin=15 * mm,
        leftMargin=15 * mm,
        topMargin=15 * mm,
        bottomMargin=15 * mm,
    )

    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        "CamerasTitle",
        parent=styles["Title"],
        fontSize=17,
        alignment=TA_CENTER,
        spaceAfter=4 * mm,
    )

    subtitle_style = ParagraphStyle(
        "CamerasSubtitle",
        parent=styles["Normal"],
        fontSize=9,
        alignment=TA_CENTER,
        spaceAfter=8 * mm,
    )

    cell_style = ParagraphStyle(
        "CamerasCell",
        parent=styles["Normal"],
        fontSize=8,
        leading=10,
    )

    cameras_db = cameras_db or []

    story = [
        Paragraph("LAPORAN MONITORING CCTV", title_style),
        Paragraph(
            f"Tanggal/Waktu Laporan: {datetime.now().strftime('%d-%m-%Y %H:%M:%S')} "
            f"— Jumlah Kamera CCTV: {len(cameras_db)}",
            subtitle_style,
        ),
    ]

    if not cameras_db:
        story.append(
            Paragraph("Tidak ada data kamera CCTV.", styles["Normal"])
        )
    else:
        data = [
            [
                Paragraph("<b>No</b>", cell_style),
                Paragraph("<b>Nama CCTV</b>", cell_style),
                Paragraph("<b>URL Stream</b>", cell_style),
                Paragraph("<b>Status</b>", cell_style),
            ]
        ]

        rows_info = []

        for index, camera in enumerate(cameras_db, start=1):
            name = _value(
                camera,
                "location",
                "name",
                "camera_name",
                default="-",
            )

            url = _value(
                camera,
                "stream_url",
                "url",
                default="-",
            )

            status = _camera_status(
                _value(
                    camera,
                    "status",
                    "active",
                    default=None,
                )
            )

            data.append(
                [
                    str(index),
                    Paragraph(_safe_text(name), cell_style),
                    Paragraph(_safe_text(url), cell_style),
                    Paragraph(status, cell_style),
                ]
            )

            rows_info.append({"row": len(data) - 1, "status": status})

        table = Table(
            data,
            colWidths=[
                12 * mm,
                45 * mm,
                95 * mm,
                28 * mm,
            ],
            repeatRows=1,
        )

        table_style = [
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2563EB")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("ALIGN", (0, 0), (-1, 0), "CENTER"),
            ("ALIGN", (0, 1), (0, -1), "CENTER"),
            ("ALIGN", (3, 1), (3, -1), "CENTER"),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#D1D5DB")),
            (
                "ROWBACKGROUNDS",
                (0, 1),
                (-1, -1),
                [colors.white, colors.HexColor("#F9FAFB")],
            ),
        ]

        for info in rows_info:
            row = info["row"]
            text = info["status"].lower()

            if "non" in text:
                fill = colors.HexColor("#F3F4F6")
            elif "error" in text:
                fill = colors.HexColor("#FEE2E2")
            elif "aktif" in text:
                fill = colors.HexColor("#DCFCE7")
            else:
                fill = colors.HexColor("#F3F4F6")

            table_style.append(
                ("BACKGROUND", (3, row), (3, row), fill)
            )

        table.setStyle(TableStyle(table_style))
        story.append(table)

    doc.build(story)

    buffer.seek(0)

    return buffer